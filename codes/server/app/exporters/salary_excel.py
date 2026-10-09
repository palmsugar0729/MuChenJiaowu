"""工资结算表 —— 从 `codes/desktop/excel_exporter.py` 移植。

★ 这个模块的目标是**与桌面版逐单元格一致**：同一批数据，桌面版导一份、
  服务端导一份，打开来每个单元格的值、公式、底色、列宽、对齐都一样。
  `docs/实施计划.md` 第六节把这条定为 Phase 4 的硬性验收标准 ——
  这张表是交给领导的，格式一个字都不能变。

  ⚠️「一致」指的是**单元格一致**，不是文件字节一致。openpyxl 每次 `save()`
     都会往 `docProps/core.xml` 写当前时间戳，同一个工作簿存两次字节都不同。
     验收方法文档里写的也是「逐单元格比对」。`tests/test_export.py` 里有一条
     测试直接把桌面版模块加载进来对拍。

⚠️ 下面的 `HEADERS` / `WEEKDAY_NAMES` / `COLOR_PALETTE` / 字体边框常量是**照抄**
   桌面版的，不是「顺便重新设计」的。出处：
     codes/desktop/excel_exporter.py   表头、样式、合并、公式、列宽
     codes/desktop/config.py:45-69     WEEKDAY_NAMES、COLOR_PALETTE
   改这里之前，先想清楚为什么 —— 特别是配色那 12 个是**打印用**的，
   跟界面的品牌色是两回事，别互相借。

⚠️ 本模块**只依赖标准库 + openpyxl**，不 import app.models / app.db。
   数据由调用方整理成 `ExportRow` 传进来 —— 这样它能被 service 用，
   也能脱离数据库直接单测。

与桌面版**有意**的差异（两处）：
  1. **费率来源**：桌面版 `Record.rate` 是查 `config.CLASS_RATES` 静态表（导出时现查）；
     这里由 `ExportRow.rate` 直接给，来源是 `lessons.rate`（排课时的快照）。
     **这样才对** —— 改一次班级费率不该篡改历史工资表。所以拿「改过费率」的月份
     和桌面版比对必然对不上，那不是 bug。
  2. **多人导出**：桌面版一次只导一位老师；这里支持一个工作簿多个 sheet（方案 A）。
     单人导出仍然沿用「工资结算表」这个 sheet 名，跟桌面版一致。
"""

import re
from dataclasses import dataclass, field
from datetime import date as Date
from datetime import time as Time
from io import BytesIO

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

# A~J 共10列（班级名不导出，仅用于配色；去掉了原来无意义的空列）
HEADERS = [
    "日期", "星期", "时间", "小时数", "时薪（元/时）",
    "薪酬", "合计时数（小时）", "合计薪酬（元）", "班级", "备注",
]

COLUMN_WIDTHS = [13, 8, 10, 10, 14, 10, 14, 14, 10, 12]

# 值班表用的是「一」不是「周一」—— 抄自 codes/desktop/config.py:45-53
WEEKDAY_NAMES: dict[int, str] = {
    0: "一",
    1: "二",
    2: "三",
    3: "四",
    4: "五",
    5: "六",
    6: "日",
}

# 同班级同色（12 色）—— 抄自 codes/desktop/config.py:56-69
COLOR_PALETTE: list[str] = [
    "FFF5B7B1",  # 浅红
    "FFAED6F1",  # 浅蓝
    "FFA9DFBF",  # 浅绿
    "FFD7BDE2",  # 浅紫
    "FFF9E79F",  # 浅黄
    "FFEDBB99",  # 浅橙
    "FFA3E4D7",  # 浅青
    "FFF1948A",  # 珊瑚红（较饱和）
    "FF85C1E9",  # 中蓝（较饱和）
    "FF82E0AA",  # 中绿（较饱和）
    "FFC39BD3",  # 中紫（较饱和）
    "FFF4D03F",  # 金黄（较饱和）
]

SHEET_TITLE = "工资结算表"

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FILL = PatternFill(start_color="FFBDD7EE", end_color="FFBDD7EE", fill_type="solid")
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=16)
NORMAL_FONT = Font(size=11)
BOLD_FONT = Font(bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


@dataclass(frozen=True)
class ExportRow:
    """导出的一行 = 一节已完成的课。

    对应桌面版的 `models.Record`，唯一的结构差异是 **`rate` 是字段而不是查表得来** ——
    服务端的费率是 `lessons.rate` 快照，见模块 docstring。
    """

    date: Date
    start_time: Time
    hours: float
    rate: float
    class_type: str
    class_name: str
    note: str = ""

    @property
    def weekday_str(self) -> str:
        return WEEKDAY_NAMES[self.date.weekday()]

    @property
    def group_key(self) -> str:
        """配色分组键。桌面版 `Record.group_key` 就是班级名 —— 保持一致。

        ⚠️ 用**班名**而不是 `class_id`：班名有唯一约束，两者一一对应，
        但只有用班名才能跟桌面版排出同一套颜色。
        """
        return self.class_name


@dataclass
class TeacherSheet:
    """一位老师一个 sheet。"""

    teacher_name: str
    rows: list[ExportRow] = field(default_factory=list)


# ── 渲染 ──────────────────────────────────────────


def render_sheet(ws, *, teacher_name: str, rows: list[ExportRow]) -> None:
    """把一位老师的课渲染成一个 sheet。**结构照抄桌面版 `export()`。**"""
    rows = sorted(rows, key=lambda r: (r.date, r.start_time))

    # 为每个 group_key 预分配颜色（按排序后的首次出现顺序）
    color_map = {
        gk: COLOR_PALETTE[idx % len(COLOR_PALETTE)]
        for idx, gk in enumerate(dict.fromkeys(r.group_key for r in rows))
    }

    # 第1行：标题
    ws.merge_cells("A1:J1")
    ws["A1"] = SHEET_TITLE
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER

    # 第2行：姓名
    ws["G2"] = "姓名："
    ws["G2"].font = NORMAL_FONT
    ws["G2"].alignment = LEFT
    ws["H2"] = teacher_name
    ws["H2"].font = NORMAL_FONT
    ws["H2"].alignment = LEFT

    # 第3行：表头
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # 数据行：按日期+时间顺序输出，同组保持同色（无需连续）
    row_num = 4
    for rec in rows:
        color = color_map[rec.group_key]
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        _write_record_row(ws, row_num, rec, fill)
        row_num += 1

    # 合计行。⚠️ 没有记录时整行都不写 —— 跟桌面版一样。
    last_data_row = row_num - 1
    if rows:
        # G 列(7)：合计时数 = SUM(D4:Dn)
        sum_hours_cell = ws.cell(row=row_num, column=7)
        sum_hours_cell.value = f"=SUM(D4:D{last_data_row})"
        sum_hours_cell.font = BOLD_FONT
        sum_hours_cell.alignment = CENTER
        sum_hours_cell.border = THIN_BORDER

        # H 列(8)：合计薪酬 = SUM(F4:Fn)
        sum_salary_cell = ws.cell(row=row_num, column=8)
        sum_salary_cell.value = f"=SUM(F4:F{last_data_row})"
        sum_salary_cell.font = BOLD_FONT
        sum_salary_cell.alignment = CENTER
        sum_salary_cell.border = THIN_BORDER

        for col_idx in range(1, 11):
            cell = ws.cell(row=row_num, column=col_idx)
            if cell.value is None:
                cell.border = THIN_BORDER

    # 列宽
    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_record_row(ws, row: int, rec: ExportRow, fill: PatternFill) -> None:
    values = [
        rec.date,                          # A: 日期（真实 date 对象，openpyxl 会设成日期格式）
        rec.weekday_str,                   # B: 星期
        rec.start_time.strftime("%H:%M"),  # C: 时间（丢掉秒）
        rec.hours,                         # D: 小时数
        rec.rate,                          # E: 时薪（★ 快照）
        None,                              # F: 薪酬（公式）
        None,                              # G: 合计时数（留空）
        None,                              # H: 合计薪酬（留空）
        rec.class_type,                    # I: 班级（★ 是**班型**，不是班名）
        rec.note,                          # J: 备注
    ]
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx)
        if val is not None:
            cell.value = val
        cell.font = NORMAL_FONT
        cell.alignment = CENTER if col_idx <= 8 else LEFT
        cell.border = THIN_BORDER
        cell.fill = fill

    # F 列(6)：薪酬公式 =D*E。★ 必须是公式而不是算好的数值 —— 文档第六节的清单里
    # 明确要求，写死了就等于把「改课时/改费率自动重算」的能力弄丢了。
    ws.cell(row=row, column=6).value = f"=D{row}*E{row}"


# ── 输出 ──────────────────────────────────────────


def build_bytes(sheets: list[TeacherSheet], *, single: bool) -> bytes:
    """渲染成 xlsx 字节。

    `single=True` → 只有一位老师，sheet 名固定为「工资结算表」（与桌面版一致）。
    `single=False` → 每位老师一个 sheet，sheet 名 = 教师姓名（方案 A）。
    """
    if not sheets:
        raise ValueError("没有数据可导出")

    wb = Workbook()
    if single:
        ws = wb.active
        ws.title = SHEET_TITLE
        render_sheet(ws, teacher_name=sheets[0].teacher_name, rows=sheets[0].rows)
    else:
        used: set[str] = set()
        for sheet in sheets:
            ws = wb.create_sheet(title=_sheet_title(sheet.teacher_name, used))
            render_sheet(ws, teacher_name=sheet.teacher_name, rows=sheet.rows)
        wb.remove(wb.active)  # 去掉 openpyxl 自建的空表

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 名字清洗 ──────────────────────────────────────
# Excel 对 sheet 名有硬约束：≤31 字符、不许含 : \ / ? * [ ]、非空、工作簿内唯一。
# 撞上就直接抛异常，所以必须在交给 openpyxl 之前洗干净。

_FORBIDDEN_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_FORBIDDEN_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]')
_MAX_SHEET_TITLE = 31


def _sheet_title(name: str, used: set[str]) -> str:
    """清洗 + 去重。同名两位老师**都保留**（第二位起加 `(2)`），绝不合并 ——
    合并会把两个人的工资算进同一张表。"""
    cleaned = _FORBIDDEN_SHEET_CHARS.sub("_", (name or "").strip()).strip("'").strip()
    if not cleaned:
        cleaned = "未命名"
    cleaned = cleaned[:_MAX_SHEET_TITLE]

    candidate, n = cleaned, 2
    while candidate.lower() in used:
        suffix = f"({n})"
        candidate = cleaned[: _MAX_SHEET_TITLE - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def safe_filename(name: str) -> str:
    """下载文件名里不能出现 `\\ / : * ? " < > |` 和控制字符。"""
    cleaned = _FORBIDDEN_FILENAME_CHARS.sub("_", name or "").strip()
    return cleaned or "salary"
